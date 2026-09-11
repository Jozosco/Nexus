#!/usr/bin/env python3
"""관세청 GW 업로드본 이후 월을 API로 이어받기 — 업로드 계열과 같은 형식의 `{국가}_API.xlsx` (A-255).

배경(승인자 지시 2026-09-11): 전 품목·국가 업로드본이 2026-07까지 덮으므로, 그 이후 월만 data.go.kr
API로 받아 **같은 폴더·같은 형식**의 동반 파일에 기록한다. 업로드본은 읽기 전용(A-184 원본 불가침).

규약:
  - 대상 쌍 = 업로드본 파일(통합본 16years·`_API`·`.partial`·DRM 제외) 1개당 1쌍. HS는 폴더에서
    산출(하위 폴더면 10자리, 루트면 6자리 — A-196: API는 10자리 hsSgn을 받는다). 국가는 업로드 파서의
    COUNTRY_CODES(파일명 → ISO2) 재사용. 매핑 실패 국가는 건너뛰고 보고한다.
  - 조회 창 = (업로드본 ∪ 기존 API 동반 파일)의 마지막 값 월 + 1 ~ 공개 완료 월(익월 15일 규칙 — A-253).
    창이 비면 API 호출 없이 종료(미공개·미입력 월은 건너뜀).
  - 동반 파일은 업로드 형식 그대로(1행 공백·2행 헤더·`1월`~`12월` 12행·5열)이며 **창 안의 월만** 기록,
    업로드본이 이미 덮는 월은 동반 파일에서 제거한다(업로드본 우선 — 파서 병합 규칙과 동일).
  - 병렬 5 · 연속 실패 회로 차단 · 시간 예산은 customs_connector 헬퍼를 재사용한다.

사용: python scripts/customs_gw_api_increment.py [--dry-run]
출력: `{국가}_API.xlsx` 갱신 + reports/market/customs_api_increment_{date}.md + stdout 요약.
"""
from __future__ import annotations

import argparse
import os
import re
import sys
import time
from concurrent.futures import ThreadPoolExecutor
from datetime import date
from pathlib import Path

import pandas as pd

_ROOT = Path(__file__).resolve().parents[1]
if str(_ROOT) not in sys.path:
    sys.path.insert(0, str(_ROOT))

from scripts.ingest_customs_gw_uploads import (GW_ROOT, _country_code, _file_kind,  # noqa: E402
                                               _hs_from_path)
from src.pipeline.connectors import customs_connector as cc  # noqa: E402

REPORT_DIR = Path("reports/market")
API_SUFFIX = "_API"
HEADER = ["무역수지(달러)", "수출액(달러)", "수출량(kg)", "수입액(달러)", "수입량(kg)"]
# API 응답 필드 → 열 순서(업로드본 열 1~5). JSON·XML 두 계열의 이름을 모두 받는다(A-152·A-158).
FIELD_CANDIDATES = {
    "무역수지(달러)": ("balPayments", "balance", "trade_balance_usd"),
    "수출액(달러)": ("expDlr", "expAmt", "exAmt", "export_fob_usd"),
    "수출량(kg)": ("expWgt", "export_weight_kg"),
    "수입액(달러)": ("impDlr", "impAmt", "imAmt", "import_cif_usd"),
    "수입량(kg)": ("impWgt", "imWgt", "wgt", "import_weight_kg"),
}
_YM_RE = re.compile(r"^(\d{4})[.\-/]?(\d{2})$")


def _months_between(start_ym: str, end_ym: str) -> list[str]:
    out, ym = [], start_ym
    while ym <= end_ym:
        out.append(ym)
        ym = cc._ym_add(ym, 1)
    return out


def _last_value_month(path: Path) -> str | None:
    """xlsx의 값이 있는 마지막 (연,월) → YYYYMM. 값 없음이면 None."""
    best = None
    try:
        xl = pd.ExcelFile(path)
        for sheet in xl.sheet_names:
            ym = re.search(r"(\d{4})", str(sheet))
            if not ym:
                continue
            year = int(ym.group(1))
            raw = xl.parse(sheet, header=None)
            for r in range(len(raw)):
                mm = re.search(r"(\d{1,2})\s*월", str(raw.iloc[r, 0]))
                if not mm:
                    continue
                vals = pd.to_numeric(raw.iloc[r, 1:6], errors="coerce")
                if vals.notna().any() and (vals.fillna(0) != 0).any():
                    cand = f"{year:04d}{int(mm.group(1)):02d}"
                    if best is None or cand > best:
                        best = cand
    except Exception as e:  # noqa: BLE001
        print(f"  [경고] {path.name} 판독 실패(비치명): {e}")
    return best


def _read_companion(path: Path) -> dict[str, list[float | None]]:
    """기존 `_API.xlsx` → {YYYYMM: [5값]}."""
    rows: dict[str, list[float | None]] = {}
    if not path.exists():
        return rows
    try:
        xl = pd.ExcelFile(path)
        for sheet in xl.sheet_names:
            ym = re.search(r"(\d{4})", str(sheet))
            if not ym:
                continue
            raw = xl.parse(sheet, header=None)
            for r in range(len(raw)):
                mm = re.search(r"(\d{1,2})\s*월", str(raw.iloc[r, 0]))
                if not mm:
                    continue
                vals = [pd.to_numeric(raw.iloc[r, c], errors="coerce") for c in range(1, 6)]
                if all(pd.isna(v) for v in vals):
                    continue
                rows[f"{int(ym.group(1)):04d}{int(mm.group(1)):02d}"] = [None if pd.isna(v) else float(v) for v in vals]
    except Exception as e:  # noqa: BLE001
        print(f"  [경고] 동반 파일 판독 실패 — 새로 작성: {path.name} ({e})")
    return rows


def _write_companion(path: Path, rows: dict[str, list[float | None]]) -> None:
    """업로드본과 동일 레이아웃(1행 공백·2행 헤더·월 12행)으로 저장. 값 없는 연도는 시트를 만들지 않음."""
    import openpyxl
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    years = sorted({ym[:4] for ym in rows})
    for y in years:
        ws = wb.create_sheet(f"{y}년")
        ws.append([None])
        ws.append([None] + HEADER)
        for m in range(1, 13):
            vals = rows.get(f"{y}{m:02d}")
            ws.append([f"{m}월"] + (vals if vals else [None] * 5))
    if not years:
        ws = wb.create_sheet("비어있음")
        ws.append([None]); ws.append([None] + HEADER)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def _items_to_months(items: list[dict]) -> dict[str, list[float | None]]:
    """API 항목 → {YYYYMM: [5값]} — 동일 월 복수 행(하위 HS)은 합산(A-085)."""
    acc: dict[str, list[float]] = {}
    for it in items:
        raw_ym = str(it.get("year") or it.get("strtYymm") or it.get("period_start") or "").strip()
        m = _YM_RE.match(raw_ym)
        if not m:
            continue
        ym = m.group(1) + m.group(2)
        vals = []
        for col in HEADER:
            v = next((it[k] for k in FIELD_CANDIDATES[col] if k in it and it[k] not in (None, "")), None)
            try:
                vals.append(float(str(v).replace(",", "")) if v is not None else 0.0)
            except ValueError:
                vals.append(0.0)
        if ym in acc:
            acc[ym] = [a + b for a, b in zip(acc[ym], vals)]
        else:
            acc[ym] = vals
    return {k: [float(x) for x in v] for k, v in acc.items()}


def _pairs() -> list[dict]:
    out, unmapped = [], []
    for p in sorted(GW_ROOT.rglob("*.xlsx")):
        if p.stem.endswith(API_SUFFIX) or ".partial" in p.name or p.stem.lower().startswith("16years"):
            continue
        if _file_kind(p) != "xlsx":
            continue
        try:
            hs, level, commodity, use_label = _hs_from_path(p)
        except ValueError:
            continue
        code = _country_code(p.stem)
        if len(code) != 2:
            unmapped.append(f"{p.parent.name}/{p.name}")
            continue
        out.append({"path": p, "hs": hs, "level": level, "commodity": commodity, "use_label": use_label,
                    "country": code, "companion": p.with_name(f"{p.stem}{API_SUFFIX}.xlsx")})
    if unmapped:
        print(f"[경고] 국가 코드 미매핑 {len(unmapped)}건 — 건너뜀: " + " | ".join(unmapped[:8]))
    return out


def run(dry_run: bool = False, today: date | None = None) -> int:
    today = today or date.today()
    key = os.environ.get("DATA_GO_KR_SERVICE_KEY", "").strip()
    pairs = _pairs()
    if not pairs:
        print("[경고] 대상 쌍 없음 — 업로드본 폴더 확인")
        return 0
    released = cc._last_released_month(today)
    # 커버리지 = 업로드본 ∪ 동반 파일의 마지막 값 월(전 쌍 최대) — 창을 닫아 이미 받은 월의 재조회를 막는다
    cov_base = cc._upload_coverage_month() or "000000"
    cov_api = max((_last_value_month(p["companion"]) or "000000" for p in pairs if p["companion"].exists()),
                  default="000000")
    cov = max(cov_base, cov_api)
    start = cc._ym_add(cov, 1) if cov != "000000" else f"{today.year:04d}01"
    print(f"[정보] 대상 {len(pairs)}쌍 · 커버리지 업로드본 {cov_base} / 동반 {cov_api} · 공개 완료 월 {released} "
          f"→ 조회 구간 {start}~{released}")
    if start > released:
        print("[정보] 조회할 신규 월 없음 — API 호출 없이 종료")
        return 0
    if dry_run:
        print(f"[정보] dry-run — {len(pairs)}쌍 × {len(_months_between(start, released))}개월 조회 예정")
        return 0
    if not key:
        print("[경고] DATA_GO_KR_SERVICE_KEY 미설정 — 조회 불가(비치명 종료)")
        return 0

    # 연도 경계는 분할(관세청 API는 1년 이내 조회만 허용 — A-101)
    windows: list[tuple[str, str]] = []
    for ym in _months_between(start, released):
        if windows and windows[-1][1][:4] == ym[:4]:
            windows[-1] = (windows[-1][0], ym)
        else:
            windows.append((ym, ym))

    def _one(pair: dict) -> tuple[dict, dict[str, list[float | None]], str]:
        if cc._budget_exceeded() or cc._circuit_open():
            return pair, {}, "skipped"
        months: dict[str, list[float | None]] = {}
        for w0, w1 in windows:
            items = cc._fetch_customs_range(key, w0, w1, pair["hs"], pair["country"])
            time.sleep(0.3)
            if items:
                cc._mark_success()
                months.update(_items_to_months(items))
        return pair, months, ("ok" if months else "empty")

    results: list[tuple[dict, dict, str]] = []
    with ThreadPoolExecutor(max_workers=int(os.environ.get("CUSTOMS_WORKERS", "5"))) as ex:
        for res in ex.map(_one, pairs):
            results.append(res)
    if cc._circuit_open():
        print(f"[경고] 연속 실패 {cc._CONSECUTIVE_FAILS}회 — data.go.kr 응답 없음으로 판단, 남은 쌍 중단(다음 런 재시도)")

    written, empty, skipped = 0, 0, 0
    lines = [f"# 관세청 API 월 증분 — {today}", "",
             f"조회 구간 {start}~{released} · 대상 {len(pairs)}쌍 · 커버리지(업로드 {cov_base}·동반 {cov_api})", "",
             "| 품목 | HS | 국가 | 상태 | 기록 월 |", "|---|---|---|---|---|"]
    for pair, months, status in results:
        existing = _read_companion(pair["companion"])
        existing = {ym: v for ym, v in existing.items() if ym > cov_base}   # 업로드본이 덮는 월은 제거
        new_months = {ym: v for ym, v in months.items() if start <= ym <= released}
        if status == "ok" and new_months:
            existing.update(new_months)
            _write_companion(pair["companion"], existing)
            written += 1
            lines.append(f"| {pair['commodity']} | {pair['hs']} | {pair['country']} | 기록 | {', '.join(sorted(new_months))} |")
        elif status == "empty":
            empty += 1
            lines.append(f"| {pair['commodity']} | {pair['hs']} | {pair['country']} | 응답 없음(무역 부재 또는 미공개) | — |")
        else:
            skipped += 1
            lines.append(f"| {pair['commodity']} | {pair['hs']} | {pair['country']} | 미시도(예산·회로 차단) | — |")
    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    rp = REPORT_DIR / f"customs_api_increment_{today.isoformat()}.md"
    rp.write_text("\n".join(lines) + "\n", encoding="utf-8")
    print(f"[완료] 동반 파일 기록 {written}쌍 · 응답 없음 {empty} · 미시도 {skipped} → {rp}")
    return 0


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--dry-run", action="store_true")
    a = ap.parse_args()
    sys.exit(run(dry_run=a.dry_run))
