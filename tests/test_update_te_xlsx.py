"""TE xlsx 증분 갱신기 — append·검증 로직 단위 테스트 (A-248, 네트워크 불요)."""
from __future__ import annotations

from datetime import date
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook

from scripts.update_te_xlsx_from_api import (
    HEADER,
    _parse_snapshot,
    _read_last_date,
    _validate,
    append_rows,
)


def _make(path: Path, desc: bool = False) -> None:
    wb = Workbook(); ws = wb.active; ws.title = "2026년"; ws.append(HEADER)
    rows = [[6, 30, 1, 2, 0.5, 1.5], [7, 1, 1, 2, 0.5, 1.6]]
    for r in (reversed(rows) if desc else rows):
        ws.append(r)
    wb.save(path)


def test_last_date_and_direction(tmp_path):
    p = tmp_path / "2010~2026_Soybeans_CBOT_USD Per Bushel.xlsx"; _make(p)
    assert _read_last_date(p) == (date(2026, 7, 1), False)
    q = tmp_path / "2019~2026_Urea_CBOT_USD Per Ton.xlsx"; _make(q, desc=True)
    assert _read_last_date(q) == (date(2026, 7, 1), True)


def test_append_ascending_and_new_year_sheet(tmp_path):
    p = tmp_path / "a.xlsx"; _make(p)
    rows = pd.DataFrame({"date": pd.to_datetime(["2026-07-02", "2027-01-04"]),
                         "open": [1, 1], "high": [2, 2], "low": [0.5, 0.5], "close": [1.7, 1.8]})
    assert append_rows(p, rows, desc=False) == 2
    wb = load_workbook(p)
    assert "2027년" in wb.sheetnames and [c.value for c in wb["2027년"][1]] == HEADER
    assert [c.value for c in wb["2026년"][4]][:2] == [7, 2]
    assert _read_last_date(p)[0] == date(2027, 1, 4)


def test_append_descending_inserts_on_top(tmp_path):
    p = tmp_path / "u.xlsx"; _make(p, desc=True)
    rows = pd.DataFrame({"date": pd.to_datetime(["2026-07-02", "2026-07-03"]),
                         "open": [1, 1], "high": [2, 2], "low": [0.5, 0.5], "close": [1.7, 1.8]})
    append_rows(p, rows, desc=True)
    ws = load_workbook(p)["2026년"]
    assert [ws.cell(2, 1).value, ws.cell(2, 2).value] == [7, 3]     # 최신이 최상단
    assert _read_last_date(p) == (date(2026, 7, 3), True)


def test_validate_rejects_regression(tmp_path):
    a = tmp_path / "a.xlsx"; _make(a)
    b = tmp_path / "b.xlsx"; wb = Workbook(); ws = wb.active; ws.title = "2026년"; ws.append(HEADER); ws.append([6, 1, 1, 1, 1, 1]); wb.save(b)
    ok, msg = _validate(a, b)
    assert not ok and "역행" in msg


# ── 스냅샷 파싱(A-250) — /markets/symbol 응답 1건 해석 ────────────────────────────
_LAST = date(2026, 9, 10)       # 수동 업로드본 기준선
_TODAY = date(2026, 9, 11)      # 금요일


def test_snapshot_minimal_datetime_last():
    """DateTime/Last만 있어도 종가 1행이 나오고, O/H/L은 NaN으로 남는다."""
    row, cols, st = _parse_snapshot(
        {"Symbol": "S 1:COM", "DateTime": "2026-09-11T21:00:00", "Last": 1042.25},
        _LAST, _TODAY)
    assert st == "ok" and row is not None
    assert row["date"] == date(2026, 9, 11) and row["close"] == 1042.25
    assert all(pd.isna(row[k]) for k in ("open", "high", "low"))
    assert "DateTime" in cols


def test_snapshot_optional_ohl_and_close_is_previous():
    """Open/DayHigh/DayLow가 오면 채우고, 종가는 Close(전일)가 아니라 Last를 쓴다."""
    row, _, st = _parse_snapshot(
        {"Date": "2026-09-11", "Last": 101.0, "Close": 99.0,
         "Open": 100.0, "DayHigh": 102.5, "DayLow": 98.5}, _LAST, _TODAY)
    assert st == "ok"
    assert (row["open"], row["high"], row["low"], row["close"]) == (100.0, 102.5, 98.5, 101.0)


def test_snapshot_rejects_weekend_future_and_stale():
    """주말 스탬프·미래 일자·파일 마지막 일자 이하는 전부 거부."""
    weekend, _, st_w = _parse_snapshot({"DateTime": "2026-09-12T10:00:00", "Last": 1.0},
                                       _LAST, date(2026, 9, 14))     # 토요일
    future, _, st_f = _parse_snapshot({"DateTime": "2026-09-14T10:00:00", "Last": 1.0},
                                      _LAST, _TODAY)
    stale, _, st_s = _parse_snapshot({"DateTime": "2026-09-10T21:00:00", "Last": 1.0},
                                     _LAST, _TODAY)
    assert (weekend, st_w) == (None, "weekend")
    assert (future, st_f) == (None, "future")
    assert (stale, st_s) == (None, "stale")


def test_snapshot_missing_fields_returns_column_list():
    """날짜·종가 컬럼이 없으면 None + 응답 컬럼 목록(로그·보고서용)."""
    row, cols, st = _parse_snapshot({"Symbol": "S 1:COM", "Importance": 1}, _LAST, _TODAY)
    assert row is None and st == "fields"
    assert cols == ["Importance", "Symbol"]
    row2, _, st2 = _parse_snapshot({"DateTime": "2026-09-11", "Last": "n/a"}, _LAST, _TODAY)
    assert row2 is None and st2 == "fields"
