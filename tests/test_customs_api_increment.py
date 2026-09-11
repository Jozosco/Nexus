"""관세청 API 월 증분 — 항목 변환·동반 파일 왕복·업로드 우선 병합·창 판정 (A-255)."""
from __future__ import annotations

from datetime import date
from pathlib import Path

import openpyxl
import pandas as pd
import pytest

from scripts import customs_gw_api_increment as inc
from scripts import ingest_customs_gw_uploads as up


def _upload_xlsx(path: Path, year: int, months: dict[int, list[float]]) -> None:
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    ws = wb.create_sheet(f"{year}년")
    ws.append([None]); ws.append([None] + inc.HEADER)
    for m in range(1, 13):
        ws.append([f"{m}월"] + (months.get(m) or [None] * 5))
    path.parent.mkdir(parents=True, exist_ok=True); wb.save(path)


def test_items_to_months_handles_both_field_families_and_sums_subcodes():
    items = [{"year": "2026.08", "balPayments": "-10", "expDlr": "0", "expWgt": "0", "impDlr": "100", "impWgt": "50"},
             {"year": "2026.08", "balPayments": "-5", "expDlr": "0", "expWgt": "0", "impDlr": "20", "impWgt": "10"},
             {"year": "2026", "impDlr": "999"},                                  # 연간 집계행 — 제외
             {"period_start": "202609", "imAmt": "7", "wgt": "3", "exAmt": "0"}]  # JSON 계열
    out = inc._items_to_months(items)
    assert out["202608"] == [-15.0, 0.0, 0.0, 120.0, 60.0]
    assert out["202609"][3] == 7.0 and out["202609"][4] == 3.0 and "2026" not in out


def test_companion_roundtrip_and_layout(tmp_path):
    p = tmp_path / "Vietnam_API.xlsx"
    inc._write_companion(p, {"202608": [1, 2, 3, 4, 5], "202609": [None, None, None, 9, 8]})
    wb = openpyxl.load_workbook(p, read_only=True)
    ws = wb["2026년"]; rows = list(ws.iter_rows(values_only=True))
    assert rows[0][0] is None and rows[1][1] == inc.HEADER[0] and rows[2][0] == "1월"   # 업로드본 레이아웃
    back = inc._read_companion(p)
    assert back["202608"] == [1, 2, 3, 4, 5] and back["202609"][3] == 9 and "202601" not in back


def test_parser_prefers_upload_over_api(tmp_path, monkeypatch):
    root = tmp_path / "GW"
    folder = root / "Soybean Oil (1507)" / "1507.90" / "Food use (.1010)"
    _upload_xlsx(folder / "Vietnam.xlsx", 2026, {7: [0, 0, 0, 100, 50]})
    inc._write_companion(folder / "Vietnam_API.xlsx", {"202607": [0, 0, 0, 999, 999], "202608": [0, 0, 0, 200, 80]})
    monkeypatch.setattr(up, "GW_ROOT", root)
    monkeypatch.setattr(up, "OUT", tmp_path / "out.parquet")
    out = up.run(today=date(2026, 10, 1))
    imp = out[(out["indicator_code"] == "KCS_1507901010_IMP_USD_VN")].set_index("price_date")["value"]
    assert imp[pd.Timestamp("2026-07-01")] == 100.0        # 업로드본 우선(999 아님)
    assert imp[pd.Timestamp("2026-08-01")] == 200.0        # 업로드본이 없는 월만 API
    kinds = out.loc[out["indicator_code"] == "KCS_1507901010_IMP_USD_VN", "source_kind"].tolist()
    assert sorted(kinds) == ["api", "upload"]


def test_run_skips_without_window(tmp_path, monkeypatch, capsys):
    root = tmp_path / "GW"
    _upload_xlsx(root / "Palm Oil (1511.10)" / "Malaysia.xlsx", 2026, {7: [0, 0, 0, 1, 1]})
    monkeypatch.setattr(up, "GW_ROOT", root); monkeypatch.setattr(inc, "GW_ROOT", root)
    monkeypatch.setenv("CUSTOMS_UPLOAD_COVERAGE_YYYYMM", "202607")
    monkeypatch.setattr(inc.cc, "_fetch_customs_range", lambda *a, **k: (_ for _ in ()).throw(AssertionError("호출 금지")))
    assert inc.run(today=date(2026, 9, 11)) == 0
    assert "조회할 신규 월 없음" in capsys.readouterr().out


def test_run_writes_companion_for_new_month(tmp_path, monkeypatch):
    root = tmp_path / "GW"
    _upload_xlsx(root / "Palm Oil (1511.10)" / "Malaysia.xlsx", 2026, {7: [0, 0, 0, 1, 1]})
    monkeypatch.setattr(up, "GW_ROOT", root); monkeypatch.setattr(inc, "GW_ROOT", root)
    monkeypatch.setattr(inc, "REPORT_DIR", tmp_path / "rep")
    monkeypatch.setenv("CUSTOMS_UPLOAD_COVERAGE_YYYYMM", "202607")
    monkeypatch.setenv("DATA_GO_KR_SERVICE_KEY", "k")
    calls = []

    def fake(key, w0, w1, hs, cc):
        calls.append((w0, w1, hs, cc))
        return [{"year": "2026.08", "balPayments": "0", "expDlr": "0", "expWgt": "0", "impDlr": "300", "impWgt": "120"}]
    monkeypatch.setattr(inc.cc, "_fetch_customs_range", fake)
    assert inc.run(today=date(2026, 9, 20)) == 0                    # 9/16 이후 → 202608 공개
    assert calls == [("202608", "202608", "151110", "MY")]
    comp = inc._read_companion(root / "Palm Oil (1511.10)" / "Malaysia_API.xlsx")
    assert comp["202608"][3] == 300.0
