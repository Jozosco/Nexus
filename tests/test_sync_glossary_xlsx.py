"""용어 사전집 xlsx ↔ entities/ontology yaml 동기화 스크립트 테스트.

실제 사전집·현행 YAML을 tmp_path에 복사해 검증한다(원본 무변경). YAML 용어 수는 시점에 따라
다르므로 단정은 상대적으로 둔다(시트 ID 집합 == YAML ID 집합).
"""
from __future__ import annotations

import re
import shutil
from pathlib import Path

import openpyxl
import pytest
import yaml

from scripts import sync_glossary_xlsx as sync

ROOT = Path(__file__).resolve().parent.parent
REAL_XLSX = ROOT / "data" / "raw" / "비정형 분석용_핵심 용어 사전집_v1.0.xlsx"
REAL_ENTITIES = ROOT / "src" / "semantic" / "entities.yaml"
REAL_ONTOLOGY = ROOT / "src" / "semantic" / "ontology.yaml"

pytestmark = pytest.mark.skipif(
    not (REAL_XLSX.exists() and REAL_ENTITIES.exists() and REAL_ONTOLOGY.exists()),
    reason="사전집 xlsx 또는 시맨틱 YAML 없음",
)


@pytest.fixture
def workset(tmp_path: Path) -> dict[str, Path]:
    """실제 xlsx·YAML을 tmp_path로 복사 — 원본은 절대 건드리지 않는다."""
    paths = {
        "xlsx": tmp_path / "glossary.xlsx",
        "entities": tmp_path / "entities.yaml",
        "ontology": tmp_path / "ontology.yaml",
    }
    shutil.copy(REAL_XLSX, paths["xlsx"])
    shutil.copy(REAL_ENTITIES, paths["entities"])
    shutil.copy(REAL_ONTOLOGY, paths["ontology"])
    return paths


@pytest.fixture
def synced(workset: dict[str, Path]) -> tuple[dict[str, Path], sync.SyncStats]:
    stats = sync.run_sync(workset["xlsx"], workset["entities"], workset["ontology"], save=True)
    return workset, stats


def _terms_by_id(entities: Path) -> dict[str, dict]:
    return {t["term_id"]: t for t in sync.load_terms(entities)}


def test_term_ids_match_yaml(synced: tuple[dict[str, Path], sync.SyncStats]) -> None:
    """시트 용어ID 집합 == YAML term_id 집합, 행 수 == 헤더 + 용어 수."""
    paths, _ = synced
    ids = set(_terms_by_id(paths["entities"]))
    ws = openpyxl.load_workbook(paths["xlsx"])[sync.SHEET_TERMS]
    sheet_ids = [ws.cell(row=r, column=1).value for r in range(2, ws.max_row + 1)]
    assert set(sheet_ids) == ids
    assert len(sheet_ids) == len(set(sheet_ids)), "용어ID 중복"
    assert ws.max_row == 1 + len(ids)


def test_term056_alert_criteria_follows_yaml(
        synced: tuple[dict[str, Path], sync.SyncStats]) -> None:
    paths, _ = synced
    expected = _terms_by_id(paths["entities"])["TERM-056"]["alert_criteria"]
    ws = openpyxl.load_workbook(paths["xlsx"])[sync.SHEET_TERMS]
    row = next(r for r in range(2, ws.max_row + 1) if ws.cell(row=r, column=1).value == "TERM-056")
    assert ws.cell(row=row, column=openpyxl.utils.column_index_from_string("Q")).value == expected


def test_appended_row_keeps_body_style(synced: tuple[dict[str, Path], sync.SyncStats]) -> None:
    paths, _ = synced
    ws = openpyxl.load_workbook(paths["xlsx"])[sync.SHEET_TERMS]
    cell = ws.cell(row=ws.max_row, column=2)
    assert cell.font.name == "Noto Sans KR"
    assert float(cell.font.size) == 10.0
    assert cell.alignment.wrap_text is True
    assert cell.alignment.vertical == "top"


def test_data_validations_preserved(synced: tuple[dict[str, Path], sync.SyncStats]) -> None:
    paths, _ = synced
    ws = openpyxl.load_workbook(paths["xlsx"])[sync.SHEET_TERMS]
    sqrefs = " ".join(str(dv.sqref) for dv in ws.data_validations.dataValidation)
    assert "R2:R500" in sqrefs
    assert "S2:S500" in sqrefs


def test_taxonomy_has_market_structure(synced: tuple[dict[str, Path], sync.SyncStats]) -> None:
    paths, _ = synced
    ws = openpyxl.load_workbook(paths["xlsx"])[sync.SHEET_TAXONOMY]
    names = {ws.cell(row=r, column=1).value for r in range(2, ws.max_row + 1)}
    assert "시장구조·기업" in names


def test_relation_codes_cover_ontology(synced: tuple[dict[str, Path], sync.SyncStats]) -> None:
    paths, _ = synced
    codes = {r["code"] for r in sync.load_relation_types(paths["ontology"])}
    ws = openpyxl.load_workbook(paths["xlsx"])[sync.SHEET_RELATIONS]
    sheet_codes = {ws.cell(row=r, column=1).value for r in range(2, ws.max_row + 1)}
    assert codes <= sheet_codes
    sym = {ws.cell(row=r, column=1).value: ws.cell(row=r, column=5).value
           for r in range(2, ws.max_row + 1)}
    assert set(sym.values()) <= {"Y", "N"}


def test_source_codes_cover_all_tokens(synced: tuple[dict[str, Path], sync.SyncStats]) -> None:
    paths, _ = synced
    tokens = set(sync.source_tokens(sync.load_terms(paths["entities"])))
    ws = openpyxl.load_workbook(paths["xlsx"])[sync.SHEET_SOURCES]
    sheet_codes = {ws.cell(row=r, column=1).value for r in range(2, ws.max_row + 1)}
    assert tokens <= sheet_codes


def test_code_list_has_status_rows(synced: tuple[dict[str, Path], sync.SyncStats]) -> None:
    paths, _ = synced
    ws = openpyxl.load_workbook(paths["xlsx"])[sync.SHEET_CODES]
    pairs = {(ws.cell(row=r, column=1).value, ws.cell(row=r, column=2).value)
             for r in range(2, ws.max_row + 1)}
    for _, value, _, _ in sync.STATUS_CODE_ROWS:
        assert ("Status", value) in pairs


def test_summary_sheet_updated(synced: tuple[dict[str, Path], sync.SyncStats]) -> None:
    paths, _ = synced
    ws = openpyxl.load_workbook(paths["xlsx"])[sync.SHEET_SUMMARY]
    assert str(ws["B6"].value).endswith("$A$500)")
    assert str(ws["B7"].value).startswith("=COUNTA(분류체계!")
    assert "$500" in str(ws["B8"].value) and "$500" in str(ws["B9"].value)
    assert "v1.1" in str(ws["A1"].value) and "v1.0" not in str(ws["A1"].value)
    assert re.search(r"\d{4}-\d{2}-\d{2}", str(ws["A3"].value)).group(0) == sync.SUMMARY_DATE
    assert ws["D14"].value == "시장구조·기업"
    assert str(ws["E14"].value).endswith("D14)")
    assert ws["D14"].font.name == ws["D13"].font.name
    assert "A1:H2" in {str(m) for m in ws.merged_cells.ranges}


def test_second_run_is_idempotent_and_check_passes(
        synced: tuple[dict[str, Path], sync.SyncStats]) -> None:
    paths, first = synced
    assert first.total > 0 or first.terms_added == 0  # 첫 실행은 보통 변경 있음(이미 동기화면 0)
    second = sync.run_sync(paths["xlsx"], paths["entities"], paths["ontology"], save=True)
    assert second.total == 0
    assert second.log == []
    rc = sync.main(["--xlsx", str(paths["xlsx"]), "--entities", str(paths["entities"]),
                    "--ontology", str(paths["ontology"]), "--check"])
    assert rc == 0


def test_check_reports_drift_on_unsynced_copy(workset: dict[str, Path]) -> None:
    """동기화 전 사본이 YAML보다 뒤처져 있으면 --check는 1을 반환하고 파일은 저장하지 않는다."""
    before = workset["xlsx"].read_bytes()
    ids_sheet = {c.value for c in openpyxl.load_workbook(workset["xlsx"])[sync.SHEET_TERMS]["A"]}
    ids_yaml = set(_terms_by_id(workset["entities"]))
    if ids_yaml <= ids_sheet:
        pytest.skip("사본이 이미 YAML과 동기화 상태 — 드리프트 검증 불가")
    rc = sync.main(["--xlsx", str(workset["xlsx"]), "--entities", str(workset["entities"]),
                    "--ontology", str(workset["ontology"]), "--check"])
    assert rc == 1
    assert workset["xlsx"].read_bytes() == before


def test_measurement_note_appended_to_rules() -> None:
    rec = {"term_id": "TERM-999", "interpretation_rules": "규칙", "measurement_note": "주석",
           "synonyms": ["a", "b"], "abbr": "", "relations": ["X", "Y"], "source_code": "S1,S2"}
    values = sync.term_to_row_values(rec)
    assert values["interpretation_rules"] == "규칙 ⋅ 측정 주석: 주석"
    assert values["synonyms"] == "a; b"
    assert values["abbr"] is None
    assert values["relations"] == "X,Y"
    assert values["source_code"] == "S1,S2"


def test_unknown_source_token_gets_stub_row(capsys: pytest.CaptureFixture[str]) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["출처코드", "유형", "제목", "연도", "반영 범위", "URL"])
    ws.append(["GH-README", "GitHub", "README", "2026", "범위", "url"])
    stats = sync.SyncStats()
    sync.sync_sources(ws, ["GH-README", "UNKNOWN_2024_x.pdf"], stats)
    assert stats.sources_added == 1
    assert ws.cell(row=3, column=1).value == "UNKNOWN_2024_x.pdf"
    assert ws.cell(row=3, column=2).value == sync.SOURCE_UNKNOWN_TYPE
    assert ws.cell(row=3, column=4).value == "2024"
    assert "[경고]" in capsys.readouterr().out
